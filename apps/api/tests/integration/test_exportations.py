"""Integration coverage for administrative lead exports and their history."""

from __future__ import annotations

from copy import deepcopy
from datetime import UTC, datetime, timedelta
from importlib import import_module
from io import BytesIO
from types import SimpleNamespace
from typing import Any

from bson import ObjectId
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from gerec_api.auth.dependencies import get_current_user
from gerec_api.auth.sessions import CurrentUser
from gerec_api.config import Settings
from gerec_api.domain.exportations import ExportationError, ExportationService
from gerec_api.infrastructure.mongo.collections import MongoCollections
from gerec_api.infrastructure.mongo.exportation_repository import MongoExportationRepository
from gerec_api.infrastructure.mongo.bootstrap import SCHEMA_VALIDATORS
from gerec_api.infrastructure.mongo.indexes import INDEXES
from gerec_api.infrastructure.mongo.migrations.runner import MIGRATIONS
from gerec_api.main import create_app


NOW = datetime(2026, 9, 22, 15, 30, tzinfo=UTC)
HEADERS = [
    "ID do lead",
    "ID da fila manual",
    "Nome",
    "Responsável atual",
    "Telefone",
    "E-mail",
    "Campanha",
    "Origem",
    "Situação comercial",
    "Desqualificado",
    "Data de criação",
    "Data de atribuição",
    "Última atualização",
    "Tipo de origem",
]


class FakeCursor(list[dict[str, Any]]):
    def sort(self, key_or_list, direction=None):
        fields = key_or_list if isinstance(key_or_list, list) else [(key_or_list, direction)]
        for field, order in reversed(fields):
            super().sort(
                key=lambda item: (item.get(field) is not None, item.get(field)),
                reverse=order < 0,
            )
        return self

    def skip(self, amount: int):
        del self[:amount]
        return self

    def limit(self, amount: int):
        del self[amount:]
        return self


class FakeCollection:
    def __init__(self) -> None:
        self.documents: list[dict[str, Any]] = []

    def find(self, query: dict[str, Any], **_: Any):
        return FakeCursor(deepcopy(item) for item in self.documents if _matches(item, query))

    def find_one(self, query: dict[str, Any], **_: Any):
        return next((deepcopy(item) for item in self.documents if _matches(item, query)), None)

    def insert_one(self, document: dict[str, Any], **_: Any):
        stored = deepcopy(document)
        stored.setdefault("_id", ObjectId())
        self.documents.append(stored)
        return SimpleNamespace(inserted_id=stored["_id"])

    def count_documents(self, query: dict[str, Any], **_: Any) -> int:
        return sum(1 for item in self.documents if _matches(item, query))


class FakeClient:
    def start_session(self):
        raise AssertionError("exportation reads and audit inserts must not require a transaction")


class FakeDatabase:
    def __init__(self) -> None:
        self.client = FakeClient()
        self.collections: dict[str, FakeCollection] = {}

    def __getitem__(self, name: str):
        return self.collections.setdefault(name, FakeCollection())

    def command(self, command: dict[str, Any]):
        assert command == {"hello": 1}
        return {"localTime": NOW}


def _matches(document: dict[str, Any], query: dict[str, Any]) -> bool:
    for key, expected in query.items():
        actual = document.get(key)
        if isinstance(expected, dict) and "$in" in expected:
            if actual not in expected["$in"]:
                return False
        elif actual != expected:
            return False
    return True


def _settings() -> Settings:
    return Settings(
        MONGODB_URI="mongodb://localhost:27017/?replicaSet=rs0",
        MONGODB_DATABASE="gerec_leads",
        APP_SECRET="exportation-route-test-secret",
    )


def _api_client(database: FakeDatabase, user: CurrentUser) -> TestClient:
    app = create_app(settings=_settings(), database=database)
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app)


def _seed_admin_and_lead(database: FakeDatabase) -> tuple[ObjectId, ObjectId]:
    admin_id = ObjectId()
    seller_id = ObjectId()
    campaign_id = ObjectId()
    database[MongoCollections.USERS].insert_one(
        {
            "_id": admin_id,
            "fullName": "Yago",
            "emailNormalized": "yago@example.test",
            "role": "admin",
        }
    )
    database[MongoCollections.USERS].insert_one(
        {
            "_id": seller_id,
            "fullName": "Renato",
            "emailNormalized": "renato@example.test",
            "role": "seller",
        }
    )
    database[MongoCollections.CAMPAIGNS].insert_one(
        {"_id": campaign_id, "displayName": "Campanha Empresas"}
    )
    lead_id = ObjectId()
    database[MongoCollections.LEADS].insert_one(
        {
            "_id": lead_id,
            "manualQueueLeadId": "MAN-00000000-0000-4000-8000-000000000001",
            "contactName": "Ana Souza",
            "assigneeId": seller_id,
            "phone": "055 11 99999-1234",
            "emailNormalized": "ana@example.test",
            "campaignId": campaign_id,
            "origin": "Indicação",
            "commercialStatus": "potential",
            "isDisqualified": False,
            "createdAt": NOW,
            "assignedAt": NOW + timedelta(minutes=5),
            "updatedAt": NOW + timedelta(minutes=10),
            "source": "manual",
        }
    )
    return admin_id, lead_id


def test_exportation_service_builds_leads_only_workbook_with_text_and_sao_paulo_dates() -> None:
    """Breaks if the workbook omits approved columns, embeds history, or loses text/timezone semantics."""
    database = FakeDatabase()
    admin_id, lead_id = _seed_admin_and_lead(database)
    service = ExportationService(MongoExportationRepository(database))

    result = service.export_leads(
        CurrentUser(str(admin_id), "yago@example.test", "admin"), {}, NOW
    )

    workbook = load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.sheetnames == ["Leads"]
    sheet = workbook["Leads"]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert [cell.value for cell in sheet[2]] == [
        str(lead_id),
        "MAN-00000000-0000-4000-8000-000000000001",
        "Ana Souza",
        "Renato",
        "055 11 99999-1234",
        "ana@example.test",
        "Campanha Empresas",
        "Indicação",
        "potential",
        "Não",
        "2026-09-22T12:30:00-03:00",
        "2026-09-22T12:35:00-03:00",
        "2026-09-22T12:40:00-03:00",
        "manual",
    ]
    for coordinate in ("A2", "B2", "E2"):
        assert sheet[coordinate].data_type == "s"
        assert sheet[coordinate].number_format == "@"
    assert result.lead_count == 1
    assert result.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert database[MongoCollections.EXPORTATIONS].documents[0]["status"] == "success"
    assert database[MongoCollections.EXPORTATIONS].documents[0]["administratorName"] == "Yago"


def test_empty_export_is_a_valid_header_only_workbook_and_records_zero() -> None:
    """Breaks if zero leads produces a corrupt/absent file or an incorrect audit count."""
    database = FakeDatabase()
    admin_id = ObjectId()
    database[MongoCollections.USERS].insert_one(
        {"_id": admin_id, "fullName": "Yago", "emailNormalized": "yago@example.test"}
    )

    result = ExportationService(MongoExportationRepository(database)).export_leads(
        CurrentUser(str(admin_id), "yago@example.test", "admin"), {}, NOW
    )

    workbook = load_workbook(BytesIO(result.content), data_only=True)
    assert workbook["Leads"].max_row == 1
    assert [cell.value for cell in workbook["Leads"][1]] == HEADERS
    assert database[MongoCollections.EXPORTATIONS].documents[0]["leadCount"] == 0
    assert database[MongoCollections.EXPORTATIONS].documents[0]["filters"] == {}


def test_exported_user_text_cannot_be_reinterpreted_as_an_excel_formula() -> None:
    """Breaks if an imported text value can execute as a formula when an admin opens the file."""
    database = FakeDatabase()
    admin_id, _ = _seed_admin_and_lead(database)
    database[MongoCollections.LEADS].documents[0]["contactName"] = "=1+1"

    result = ExportationService(MongoExportationRepository(database)).export_leads(
        CurrentUser(str(admin_id), "yago@example.test", "admin"), {}, NOW
    )

    sheet = load_workbook(BytesIO(result.content), data_only=False)["Leads"]
    assert sheet["C2"].value == "=1+1"
    assert sheet["C2"].data_type == "s"


def test_automatic_export_prefers_original_source_phone_over_normalized_digits() -> None:
    """Breaks if an automatic lead loses its original phone formatting during export."""
    database = FakeDatabase()
    admin_id, lead_id = _seed_admin_and_lead(database)
    lead = database[MongoCollections.LEADS].documents[0]
    lead.pop("phone")
    lead.pop("manualQueueLeadId")
    lead["source"] = "google_sheets"
    lead["phoneNormalized"] = "0551199991234"
    database[MongoCollections.SOURCE_RECORDS].insert_one(
        {
            "leadId": lead_id,
            "sellerProjection": {"phone_number": "(055) 11 9999-1234"},
        }
    )

    result = ExportationService(MongoExportationRepository(database)).export_leads(
        CurrentUser(str(admin_id), "yago@example.test", "admin"), {}, NOW
    )

    sheet = load_workbook(BytesIO(result.content), data_only=True)["Leads"]
    assert sheet["E2"].value == "(055) 11 9999-1234"
    assert sheet["E2"].data_type == "s"


def test_admin_download_and_paginated_history_use_the_public_http_contract() -> None:
    """Breaks if the API does not stream XLSX bytes or history pagination/order drifts."""
    database = FakeDatabase()
    admin_id, _ = _seed_admin_and_lead(database)
    history = database[MongoCollections.EXPORTATIONS]
    for offset in range(3):
        history.insert_one(
            {
                "actorId": admin_id,
                "administratorName": "Yago",
                "leadCount": offset,
                "filters": {},
                "status": "success",
                "createdAt": NOW + timedelta(minutes=offset),
            }
        )
    client = _api_client(
        database, CurrentUser(str(admin_id), "yago@example.test", "admin")
    )

    response = client.get("/api/admin/exportations/leads")
    page = client.get("/api/admin/exportations?page=2&limit=2")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")
    assert page.status_code == 200
    assert page.json() == {
        "items": [
            {
                "createdAt": NOW.isoformat().replace("+00:00", "Z"),
                "administratorName": "Yago",
                "leadCount": 0,
                "filters": {},
                "status": "success",
            },
            {
                "createdAt": NOW.isoformat().replace("+00:00", "Z"),
                "administratorName": "Yago",
                "leadCount": 1,
                "filters": {},
                "status": "success",
            },
        ],
        "page": 2,
        "pageSize": 2,
        "total": 4,
    }


def test_seller_cannot_download_or_read_exportation_history() -> None:
    """Breaks if hiding the frontend tab becomes the only export authorization control."""
    database = FakeDatabase()
    seller = CurrentUser(str(ObjectId()), "seller@example.test", "seller")
    client = _api_client(database, seller)

    assert client.get("/api/admin/exportations/leads").status_code == 403
    assert client.get("/api/admin/exportations").status_code == 403
    assert database[MongoCollections.EXPORTATIONS].documents == []


class FailingReadRepository:
    def __init__(self) -> None:
        self.records: list[dict[str, Any]] = []

    def list_leads(self, filters: dict[str, Any]) -> list[dict[str, Any]]:
        raise RuntimeError("mongodb://secret-host/private-data")

    def record_exportation(self, **record: Any) -> None:
        self.records.append(record)

    def list_exportations(self, *, page: int, limit: int) -> dict[str, Any]:
        return {"items": [], "page": page, "pageSize": limit, "total": 0}


def test_export_failure_records_only_error_and_exposes_no_internal_details() -> None:
    """Breaks if storage/generation failure is audited as success or leaks secrets to callers."""
    repository = FailingReadRepository()
    service = ExportationService(repository)

    try:
        service.export_leads(CurrentUser("admin-yago", "yago@example.test", "admin"), {}, NOW)
    except ExportationError as error:
        assert str(error) == "Não foi possível gerar a exportação."
    else:
        raise AssertionError("the export was expected to fail")

    assert len(repository.records) == 1
    assert repository.records[0]["status"] == "error"
    assert repository.records[0]["lead_count"] == 0
    assert "secret-host" not in repr(repository.records[0])


def test_http_export_failure_returns_recoverable_safe_error_without_false_success() -> None:
    """Breaks if an export failure becomes a 200/download or exposes an exception string."""
    database = FakeDatabase()
    app = create_app(settings=_settings(), database=database)
    app.state.exportation_service = ExportationService(FailingReadRepository())
    app.dependency_overrides[get_current_user] = lambda: CurrentUser(
        "admin-yago", "yago@example.test", "admin"
    )

    response = TestClient(app).get("/api/admin/exportations/leads")

    assert response.status_code == 503
    assert response.json() == {"detail": "Não foi possível gerar a exportação."}
    assert "secret-host" not in response.text


class IndexTrackingCollection:
    def __init__(self) -> None:
        self.indexes: dict[str, dict[str, Any]] = {}

    def index_information(self, **_: Any) -> dict[str, dict[str, Any]]:
        return deepcopy(self.indexes)

    def create_index(self, keys, **options: Any) -> str:
        name = options.pop("name")
        self.indexes[name] = {"key": list(keys), **options}
        return name


class MigrationDatabase:
    def __init__(self) -> None:
        self.collection = IndexTrackingCollection()

    def __getitem__(self, name: str) -> IndexTrackingCollection:
        assert name == MongoCollections.EXPORTATIONS
        return self.collection


def test_exportation_history_collection_and_indexes_are_versioned_idempotently() -> None:
    """Breaks if history deployment lacks its collection/index migration or repeats DDL unsafely."""
    migration = import_module(
        "gerec_api.infrastructure.mongo.migrations.20260923_exportation_history"
    )
    database = MigrationDatabase()

    migration.apply(database)
    migration.apply(database)

    assert MongoCollections.EXPORTATIONS in MongoCollections.ALL
    assert migration.VERSION in {version for version, _ in MIGRATIONS}
    assert {
        definition.name
        for definition in INDEXES
        if definition.collection_name == MongoCollections.EXPORTATIONS
    } == {
        "exportations_created_at_desc",
        "exportations_actor_created_at_desc",
    }
    assert set(database.collection.indexes) == {
        "exportations_created_at_desc",
        "exportations_actor_created_at_desc",
    }
    schema = SCHEMA_VALIDATORS[MongoCollections.EXPORTATIONS]["$jsonSchema"]
    assert schema["required"] == [
        "actorId",
        "administratorName",
        "leadCount",
        "filters",
        "status",
        "createdAt",
    ]
    assert schema["properties"]["status"] == {"enum": ["success", "error"]}
