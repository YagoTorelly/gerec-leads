"""Administrative lead export use case and XLSX generation."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any, Mapping, Protocol
from zoneinfo import ZoneInfo

from openpyxl import Workbook


SAO_PAULO = ZoneInfo("America/Sao_Paulo")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_HEADERS = (
    "ID do lead",
    "ID da fila manual",
    "Nome",
    "Responsável atual",
    "Telefone",
    "E-mail",
    "Campanha",
    "Origem",
    "Situação comercial",
    "Desqualificado",
    "Data de criação",
    "Data de atribuição",
    "Última atualização",
    "Tipo de origem",
)


class ExportationError(RuntimeError):
    """Safe, recoverable failure exposed by the export application service."""


class ExportationPermissionError(PermissionError):
    """Raised when a non-administrator calls the service directly."""


@dataclass(frozen=True)
class ExportationResult:
    content: bytes
    filename: str
    lead_count: int
    media_type: str = XLSX_MEDIA_TYPE


class ExportationPersistence(Protocol):
    def list_leads(self, filters: dict[str, Any]) -> list[dict[str, Any]]: ...

    def record_exportation(
        self,
        *,
        actor_id: str,
        administrator_fallback: str,
        lead_count: int,
        filters: dict[str, Any],
        status: str,
        created_at: datetime,
        error_code: str | None = None,
    ) -> None: ...

    def list_exportations(self, *, page: int, limit: int) -> dict[str, Any]: ...


class ExportationService:
    """Generate lead-only workbooks and audit exactly the attempted outcome."""

    def __init__(self, persistence: ExportationPersistence) -> None:
        self._persistence = persistence

    def export_leads(
        self,
        actor: Any,
        filters: Mapping[str, Any] | None,
        now: datetime,
    ) -> ExportationResult:
        actor_id, actor_fallback = _administrator(actor)
        timestamp = _aware_utc(now)
        normalized_filters = dict(filters or {})
        lead_count = 0
        try:
            leads = self._persistence.list_leads(normalized_filters)
            lead_count = len(leads)
            content = _workbook_bytes(leads)
        except Exception as error:
            self._record_error(
                actor_id=actor_id,
                actor_fallback=actor_fallback,
                lead_count=lead_count,
                filters=normalized_filters,
                created_at=timestamp,
            )
            raise ExportationError("Não foi possível gerar a exportação.") from error

        try:
            self._persistence.record_exportation(
                actor_id=actor_id,
                administrator_fallback=actor_fallback,
                lead_count=lead_count,
                filters=normalized_filters,
                status="success",
                created_at=timestamp,
            )
        except Exception as error:
            self._record_error(
                actor_id=actor_id,
                actor_fallback=actor_fallback,
                lead_count=lead_count,
                filters=normalized_filters,
                created_at=timestamp,
            )
            raise ExportationError("Não foi possível gerar a exportação.") from error

        local_timestamp = timestamp.astimezone(SAO_PAULO)
        return ExportationResult(
            content=content,
            filename=f"leads-{local_timestamp:%Y%m%d-%H%M%S}.xlsx",
            lead_count=lead_count,
        )

    def history(
        self,
        actor: Any,
        *,
        page: int,
        limit: int,
    ) -> dict[str, Any]:
        _administrator(actor)
        if page < 1:
            raise ValueError("page must be at least 1")
        if limit < 1 or limit > 200:
            raise ValueError("limit must be between 1 and 200")
        try:
            return self._persistence.list_exportations(page=page, limit=limit)
        except Exception as error:
            raise ExportationError("Não foi possível consultar as exportações.") from error

    def _record_error(
        self,
        *,
        actor_id: str,
        actor_fallback: str,
        lead_count: int,
        filters: dict[str, Any],
        created_at: datetime,
    ) -> None:
        try:
            self._persistence.record_exportation(
                actor_id=actor_id,
                administrator_fallback=actor_fallback,
                lead_count=lead_count,
                filters=filters,
                status="error",
                created_at=created_at,
                error_code="export_failed",
            )
        except Exception:
            # The public error remains safe even when the audit store itself is unavailable.
            return


def _workbook_bytes(leads: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    sheet.append(EXPORT_HEADERS)
    for lead in leads:
        sheet.append(
            (
                _text(lead.get("leadId")),
                _text(lead.get("manualQueueLeadId")),
                _text(lead.get("name")),
                _text(lead.get("assigneeName")),
                _text(lead.get("phone")),
                _text(lead.get("email")),
                _text(lead.get("campaign")),
                _text(lead.get("origin")),
                _text(lead.get("commercialStatus")),
                "Sim" if bool(lead.get("isDisqualified")) else "Não",
                _local_datetime(lead.get("createdAt")),
                _local_datetime(lead.get("assignedAt")),
                _local_datetime(lead.get("updatedAt")),
                _text(lead.get("sourceType")),
            )
        )
    for row in range(2, sheet.max_row + 1):
        for cell in sheet[row]:
            cell.data_type = "s"
        for column in (1, 2, 5):
            sheet.cell(row=row, column=column).number_format = "@"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _administrator(actor: Any) -> tuple[str, str]:
    role = actor.get("role") if isinstance(actor, Mapping) else getattr(actor, "role", None)
    if role != "admin":
        raise ExportationPermissionError("administrator role is required")
    actor_id = actor.get("id") if isinstance(actor, Mapping) else getattr(actor, "id", None)
    fallback = actor.get("email") if isinstance(actor, Mapping) else getattr(actor, "email", None)
    if actor_id is None or not str(actor_id).strip():
        raise ValueError("actor is required")
    return str(actor_id).strip(), _text(fallback) or str(actor_id).strip()


def _aware_utc(value: datetime) -> datetime:
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("now must be timezone-aware")
    return value.astimezone(UTC)


def _local_datetime(value: Any) -> str:
    if not isinstance(value, datetime):
        return ""
    aware = value.replace(tzinfo=UTC) if value.tzinfo is None else value
    return aware.astimezone(SAO_PAULO).isoformat(timespec="seconds")


def _text(value: Any) -> str:
    return "" if value is None else str(value)
