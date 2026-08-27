"""Read-only adapter for the provisional A-Q workbook contract."""

from collections.abc import Iterable
from pathlib import Path

from openpyxl import load_workbook

from gerec_api.domain.normalization import NormalizedSourceRow, normalize_source_row


EXPECTED_HEADERS = (
    "id",
    "created_time",
    "ad_id",
    "ad_name",
    "adset_id",
    "adset_name",
    "campaign_id",
    "campaign_name",
    "form_id",
    "form_name",
    "is_organic",
    "platform",
    "voc\u00ea_tem_cnpj_ou_mei?",
    "full_name",
    "phone_number",
    "email",
    "lead_status",
)


class InvalidWorkbookError(ValueError):
    """The workbook cannot safely represent a complete source snapshot."""


class WorkbookAdapter:
    """Validate the physical mock once, then yield normalized rows without writing it."""

    def read(self, path: Path) -> Iterable[NormalizedSourceRow]:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
        try:
            if "Leads" not in workbook.sheetnames:
                raise InvalidWorkbookError("workbook must contain the Leads sheet")
            sheet = workbook["Leads"]
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = tuple(next(rows))
            except StopIteration as error:
                raise InvalidWorkbookError("workbook has no headers") from error
            if headers != EXPECTED_HEADERS:
                raise InvalidWorkbookError("workbook headers do not match the exact A-Q contract")
            for values in rows:
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                yield normalize_source_row(dict(zip(EXPECTED_HEADERS, values, strict=True)))
        finally:
            workbook.close()
